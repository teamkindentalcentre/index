#!/usr/bin/env python3
"""One-off dev utility: parses the "Blue room check list" and "Green Room
Check List" tabs of the source spreadsheet into data/seed-items.json, which
the app uses to seed its database on first run. Not used at app runtime.

Usage: python3 scripts/extract-seed.py <path-to-xlsx>
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

SKIP_NOISE_ITEM_TEXT = {"items"}
TABLE_TOP_ALIASES = {"table top"}
XRAY_ALIASES = {"in x-ray room"}
CONSUMABLE_ALIASES = {"consumable", "consumables"}


def format_qty(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    text = str(value).strip()
    return text or None


def clean(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def clean_section_name(text):
    # Trim trailing "-"/"–" left over from labels like "L17 - " with no suffix.
    stripped = re.sub(r"[\s\-–—]+$", "", text)
    return stripped or text


def extract_room(ws, category_col, item_col, qty_col, max_row, room_id, room_name):
    sections = []  # list of {"name": str, "items": [...]}
    section_index = {}
    last_primary_section = None
    current_section_name = None

    def get_or_create_section(name):
        if name not in section_index:
            section_index[name] = len(sections)
            sections.append({"name": name, "items": []})
        return sections[section_index[name]]

    for row in range(1, max_row + 1):
        category = clean(ws.cell(row=row, column=category_col).value)
        item = clean(ws.cell(row=row, column=item_col).value)
        qty_raw = ws.cell(row=row, column=qty_col).value

        item_l = item.casefold() if item else None

        # Noise: repeated table header rows ("ITEMS" / "QTY" column titles)
        if item_l in SKIP_NOISE_ITEM_TEXT:
            continue
        # Noise: "BLUE/GREEN ROOM INSTRUMENT CHECKLIST" title rows
        if item_l and "instrument checklist" in item_l:
            continue
        # Noise: "** If there is any missing..." notices
        if (item and item.startswith("**")) or (category and category.startswith("**")):
            continue
        # Noise: rubbish bin placeholder row of slashes
        if item and set(item) <= {"/"}:
            continue

        if item_l in TABLE_TOP_ALIASES:
            last_primary_section = "Table Top"
            current_section_name = "Table Top"
            continue
        if item_l in XRAY_ALIASES:
            last_primary_section = "X-ray Room"
            current_section_name = "X-ray Room"
            continue

        if category:
            category = clean_section_name(category)
            if category.casefold() in CONSUMABLE_ALIASES:
                base = last_primary_section or "Consumable"
                current_section_name = f"{base} – Consumable"
            else:
                last_primary_section = category
                current_section_name = category
            if item:
                get_or_create_section(current_section_name)["items"].append(
                    {"name": item, "expectedQty": format_qty(qty_raw)}
                )
            continue

        if item:
            if current_section_name is None:
                # Shouldn't happen in practice, but guard just in case.
                current_section_name = "Other"
            get_or_create_section(current_section_name)["items"].append(
                {"name": item, "expectedQty": format_qty(qty_raw)}
            )

    sections = [s for s in sections if s["items"]]
    return {"id": room_id, "name": room_name, "sections": sections}


def main():
    if len(sys.argv) != 2:
        print("Usage: python3 scripts/extract-seed.py <path-to-xlsx>", file=sys.stderr)
        sys.exit(1)

    src = Path(sys.argv[1])
    wb = openpyxl.load_workbook(src, data_only=True)

    blue = extract_room(
        wb["Blue room check list"],
        category_col=2,  # B
        item_col=3,  # C
        qty_col=5,  # E
        max_row=389,
        room_id="blue",
        room_name="Blue Room",
    )
    green = extract_room(
        wb["Green Room Check List "],
        category_col=3,  # C
        item_col=4,  # D
        qty_col=6,  # F
        max_row=304,
        room_id="green",
        room_name="Green Room",
    )

    out = {"rooms": [blue, green]}

    out_path = Path(__file__).resolve().parent.parent / "data" / "seed-items.json"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(out, indent=2) + "\n")

    for room in out["rooms"]:
        total_items = sum(len(s["items"]) for s in room["sections"])
        print(f"{room['name']}: {len(room['sections'])} sections, {total_items} items")
        for s in room["sections"]:
            print(f"  - {s['name']} ({len(s['items'])})")


if __name__ == "__main__":
    main()
